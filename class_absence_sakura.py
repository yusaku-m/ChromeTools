"""一部の授業だけ欠課した学生に、学務係への連絡を促すさくら連絡網を1人ずつ送る準備をする。

    pixi run absence             # 前回の授業日の分(行事予定から求める)
    pixi run absence 1026        # 日付指定(MMDD、年は今年。YYYY-MM-DD も可)
    pixi run absence --inspect   # 出力ファイルの構造(見出し行と件数)だけ表示して終わる
    pixi run absence --show      # KOSEN WEB の取得もブラウザ画面を出して行う

KOSEN WEB からの取得はヘッドレスで行う。ログインしていなければ(パスワードの自動入力も無ければ)
画面ありで起動し直してログインを待つ。さくら連絡網は送信確認を目で見るので常に画面ありで開く。

1. 行事予定(data/*行事予定*.xlsx)から前回の授業日を求める。対応済みの日なら終了
2. KOSEN WEB「日別出欠の閲覧（期間指定）」でその日の 3年・機械工学科 を出力
3. 全時限欠課(=丸一日欠席。SHR 側で連絡済み)と欠課なしを除き、残りの学生ごとに欠課した時限を集める
4. 学生ごとにさくら連絡網のメッセージを新規作成し(受信者氏名を差し込む)、送信確認画面で止まる。
   **送信ボタンはユーザーが押す。** Enter で次の学生へ進む。

出力ファイルと氏名は個人情報なので、画面にもログにも出さない(件数・出席番号・時限だけ出す)。
出力ファイルは一時フォルダに保存し、読み終えたらすぐ消す。

TODO(docs/backlog_garoon.md の4): Garoon の「学生欠席等連絡」に欠席連絡がある学生を除外する。
"""

import argparse
import csv
import datetime
import glob
import json
import os
import re
import shutil
import sys
import tempfile
import unicodedata

import openpyxl

from Chrome.KosenWeb import KosenWeb, LoginRequired
from Chrome.Sakura import Sakura

GRADE = '3年'
DEPARTMENT = '機械工学科'
AFFILIATION = '機械工学科３年'
SUBJECT_PREFIX = '授業の欠席について'
WEEKDAYS = '月火水木金土日'
BODY = """機械工学科３年生担任の前田です。

{date}の{lessons}を欠席されていたようですが，学務係へは連絡が無かったようです。
欠席がお分かりの場合は，必ず学務係への事前連絡をお願いいたします。
以上よろしくお願いいたします。"""

#出力ファイルの列位置(ユーザー指定)。B列=氏名、I列以降=その日の各時限(最大8列)、2行目=授業名
NAME_COL = 1
FIRST_LESSON_COL = 8
MAX_LESSONS = 8
LESSON_NAME_ROW = 1
ABSENT_MARK = '欠'

SCHEDULE_GLOB = os.path.join('data', '*行事予定*.xlsx')
DONE_FILE = os.path.join('data', 'class_absence_done.json')
USER_DATA_PATH = os.path.join(os.path.expanduser('~'), 'AppData', 'Local', 'Google', 'Chrome', 'AutoSyncData')


# ----------------------------------------------------------------------
# 行事予定
# ----------------------------------------------------------------------
def load_class_days(path):
    """行事予定 xlsx から高松キャンパスの授業日の集合を返す。

    シートは月ごとのブロックが横に並ぶ(2026-10 の「令和8年度行事予定案15」で確認)。
    2行目に「4月」等の月見出しがありその列が「曜」、3行目でその右にある「高松」列の
    さらに右隣が授業週の番号。授業日なら数字、試験・行事日は「〇」、休日は空欄。
    日付 d の行は d+3 行目。
    """
    m = re.search(r'令和(\d+)年度', os.path.basename(path))
    if not m:
        raise ValueError(f"ファイル名から年度を読めません: {path}")
    nendo = 2018 + int(m.group(1))
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets[0]
    rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=34, values_only=True)]

    days = set()
    for col, header in enumerate(rows[1]):
        mm = re.fullmatch(r'(\d+)月', str(header or '').strip())
        if not mm:
            continue
        month = int(mm.group(1))
        year = nendo if month >= 4 else nendo + 1
        takamatsu = [c for c in range(col + 1, col + 6) if rows[2][c] == '高松']
        if len(takamatsu) != 1:
            raise ValueError(f"{month}月の「高松」列が見つかりません")
        week_col = takamatsu[0] + 1
        for d in range(1, 32):
            try:
                date = datetime.date(year, month, d)
            except ValueError:
                break
            row = rows[d + 2]
            if row[col] != WEEKDAYS[date.weekday()]:
                raise ValueError(f"行事予定の {date} の曜日が合いません({row[col]})。シートの形が変わった可能性があります")
            if isinstance(row[week_col], (int, float)):
                days.add(date)
    return days


def previous_class_day(today, class_days):
    earlier = [d for d in class_days if d < today]
    if not earlier:
        raise ValueError(f"{today} より前の授業日が行事予定にありません")
    return max(earlier)


def find_schedule_file():
    files = glob.glob(SCHEDULE_GLOB)
    if not files:
        sys.exit(f"行事予定が見つかりません({SCHEDULE_GLOB})")
    return max(files, key=os.path.getmtime)


# ----------------------------------------------------------------------
# 対応済みの日
# ----------------------------------------------------------------------
def load_done():
    try:
        with open(DONE_FILE, encoding='utf-8') as f:
            return set(json.load(f))
    except FileNotFoundError:
        return set()


def mark_done(date):
    done = load_done() | {date.isoformat()}
    with open(DONE_FILE, 'w', encoding='utf-8') as f:
        json.dump(sorted(done), f)


# ----------------------------------------------------------------------
# KOSEN WEB の出力ファイル(個人情報。中身は表示しない)
# ----------------------------------------------------------------------
def read_table(path):
    """CSV/xlsx を文字列の2次元リストで読む。"""
    if path.lower().endswith(('.xlsx', '.xlsm')):
        #read_only=True だとファイルに書かれた <dimension> の範囲しか読まない。KOSEN WEB の出力は
        #それが実データより狭く、右側の列が欠けるので通常モードで読む(ファイルは小さい)
        ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
        return [['' if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
    for encoding in ('utf-8-sig', 'cp932'):
        try:
            with open(path, encoding=encoding, newline='') as f:
                return list(csv.reader(f))
        except UnicodeDecodeError:
            continue
    raise ValueError(f"文字コードを判別できません: {os.path.basename(path)}")


def _cell(row, col):
    return unicodedata.normalize('NFKC', row[col]).strip() if col < len(row) and row[col] else ''


def lesson_columns(rows):
    """[(列, 時限, 授業名), ...]。授業名(2行目)が空の列は含めない。

    時限は1行目の見出しに「N限」があればそれを使い、無ければ I 列=1限として列位置から決める。
    """
    result = []
    for i in range(MAX_LESSONS):
        col = FIRST_LESSON_COL + i
        name = _cell(rows[LESSON_NAME_ROW], col)
        if not name:
            continue
        m = re.search(r'(\d+)\s*限', _cell(rows[0], col) + ' ' + name)
        result.append((col, int(m.group(1)) if m else i + 1, name))
    return result


def find_targets(rows, lessons):
    """一部の時限だけ欠課した学生を [(行番号, 出席番号らしき値, 氏名, [(時限, 授業名), ...]), ...] で返す。
    あわせて (全時限欠課の人数, 欠課なしの人数) を返す。"""
    targets, all_absent, none_absent = [], 0, 0
    for r, row in enumerate(rows[LESSON_NAME_ROW + 1:], start=LESSON_NAME_ROW + 2):
        name = _cell(row, NAME_COL)
        if not name:
            continue
        absent = [(period, lesson) for col, period, lesson in lessons if _cell(row, col) == ABSENT_MARK]
        if not absent:
            none_absent += 1
        elif len(absent) == len(lessons):
            all_absent += 1
        else:
            targets.append((r, _cell(row, 0), name, absent))
    return targets, all_absent, none_absent


def inspect(rows, lessons):
    """出力ファイルの構造だけを表示する(氏名や個々の学生の値は出さない)。"""
    print(f"行数: {len(rows)} / 列数: {max((len(r) for r in rows), default=0)}")
    for i in range(min(2, len(rows))):
        print(f"{i + 1}行目: {rows[i]}")
    print("授業の列:", [(chr(ord('A') + c), p, n) for c, p, n in lessons])
    marks = {}
    for row in rows[LESSON_NAME_ROW + 1:]:
        if _cell(row, NAME_COL):
            for col, _, _ in lessons:
                v = _cell(row, col)
                marks[v] = marks.get(v, 0) + 1
    print("授業の列に現れる値と件数:", marks)


def format_lessons(absent):
    """[(1,'A'),(2,'A'),(4,'B')] → '1・2限 A，4限 B'。同じ授業の連続した時限はまとめる。"""
    groups = []
    for period, lesson in absent:
        if groups and groups[-1][1] == lesson and groups[-1][0][-1] == period - 1:
            groups[-1][0].append(period)
        else:
            groups.append(([period], lesson))
    return '，'.join('・'.join(map(str, periods)) + '限 ' + lesson for periods, lesson in groups)


def format_date(date):
    return f'{date.month}/{date.day}({WEEKDAYS[date.weekday()]})'


def parse_date(text):
    """'1026' (MMDD、年は今年) または '2026-10-26' を date にする。"""
    try:
        if len(text) == 4 and text.isdigit():
            return datetime.date(datetime.date.today().year, int(text[:2]), int(text[2:]))
        return datetime.date.fromisoformat(text)
    except ValueError:
        raise argparse.ArgumentTypeError(f"日付は MMDD か YYYY-MM-DD で指定してください: {text}")


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument('date', nargs='?', type=parse_date,
                        help='対象日 MMDD(今年) または YYYY-MM-DD。省略時は行事予定から求めた前回の授業日')
    parser.add_argument('--force', action='store_true', help='対応済みの日でも実行する')
    parser.add_argument('--inspect', action='store_true', help='出力ファイルの構造だけ表示して終わる')
    parser.add_argument('--list-days', action='store_true', help='行事予定から求めた授業日の一覧を表示して終わる')
    parser.add_argument('--show', action='store_true', help='KOSEN WEB の取得もヘッドレスにせず画面を出して行う')
    parser.add_argument('--debugger-address', help='起動済みChromeに接続する(例 127.0.0.1:9333)')
    args = parser.parse_args()

    if args.list_days:
        schedule = find_schedule_file()
        days = sorted(load_class_days(schedule))
        print(f"行事予定: {os.path.basename(schedule)}  授業日 {len(days)} 日")
        for month in sorted({(d.year, d.month) for d in days}):
            in_month = [d for d in days if (d.year, d.month) == month]
            print(f"{month[0]}/{month[1]:02d} ({len(in_month)}日): "
                  + ' '.join(f"{d.day}{WEEKDAYS[d.weekday()]}" for d in in_month))
        return

    date = args.date
    if date is None:
        schedule = find_schedule_file()
        date = previous_class_day(datetime.date.today(), load_class_days(schedule))
        print(f"行事予定: {os.path.basename(schedule)}")
    print(f"対象日: {date} ({WEEKDAYS[date.weekday()]})")
    if date.isoformat() in load_done() and not (args.force or args.inspect):
        print("この日は対応済みです(やり直すときは --force)。")
        return

    kosen = KosenWeb(USER_DATA_PATH, debugger_address=args.debugger_address, headless=not args.show)
    download_dir = tempfile.mkdtemp(prefix='kosenweb_')
    try:
        try:
            path = kosen.download_daily_attendance(date, GRADE, DEPARTMENT, download_dir)
        except LoginRequired:
            print("KOSEN WEB にログインしていないので、ブラウザ画面を出して起動し直します。")
            kosen.driver.quit()
            kosen = KosenWeb(USER_DATA_PATH, debugger_address=args.debugger_address)
            path = kosen.download_daily_attendance(date, GRADE, DEPARTMENT, download_dir)
        rows = read_table(path)
    finally:
        shutil.rmtree(download_dir, ignore_errors=True)

    if kosen.headless:
        #画面の無い Chrome はもう使わない(さくら連絡網は画面ありで開き直す)
        kosen.driver.quit()
    lessons = lesson_columns(rows) if len(rows) > LESSON_NAME_ROW else []
    if args.inspect:
        inspect(rows, lessons)
        return
    if not lessons:
        sys.exit("出力ファイルに授業の列がありません。--inspect で構造を確認してください。")
    targets, all_absent, none_absent = find_targets(rows, lessons)
    print(f"授業 {len(lessons)} コマ / 全時限欠課 {all_absent} 名(除外) / 欠課なし {none_absent} 名(除外) / 対象 {len(targets)} 名")
    if not targets:
        mark_done(date)
        print("送信対象がいないので終了します。")
        return

    if kosen.headless:
        #送信内容を目で確認するので、さくら連絡網は画面ありの Chrome で開く
        sakura = Sakura(USER_DATA_PATH)
    else:
        sakura = Sakura.share_driver(kosen)
        sakura.login()
    subject = f'{SUBJECT_PREFIX}【{date.month}/{date.day}（{WEEKDAYS[date.weekday()]}）】'
    skipped = 0
    for i, (_, number, name, absent) in enumerate(targets, start=1):
        lessons_text = format_lessons(absent)
        print(f"[{i}/{len(targets)}] 出席番号 {number}: {lessons_text}")
        body = BODY.format(date=format_date(date), lessons=lessons_text)
        sakura.new_message(subject, body)
        unmatched = sakura.set_recipients([name], AFFILIATION)
        if unmatched:
            print("  さくら連絡網で宛先が見つかりませんでした。この学生は手動で対応してください。")
            skipped += 1
            continue
        for text in sakura.go_send_detail():
            print("  ", text)
        for text in sakura.set_insert_recipient_name():
            print("  ", text)
        answer = input("  ブラウザで内容を確認して送信してください。送信したら Enter / 送らない場合は s / 中断は q: ").strip().lower()
        if answer == 'q':
            print("中断しました(この日は対応済みにしません)。")
            return
        if answer == 's':
            skipped += 1

    mark_done(date)
    print(f"完了: {len(targets) - skipped} 名に送信、{skipped} 名は送信せず。{date} を対応済みにしました。")


if __name__ == '__main__':
    main()
